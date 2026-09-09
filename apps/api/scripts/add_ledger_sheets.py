#!/usr/bin/env python3
"""Add the three whole-state ledger sheets to দৈনিক খরচের হিসাব.

    uv --directory apps/api run --with openpyxl \
        python scripts/add_ledger_sheets.py in.xlsx out.xlsx

`--with openpyxl` rather than a dependency: this runs when the workbook
template changes, which is roughly never, and the API image has no reason to
carry a spreadsheet library it never imports.

WHY THIS EXISTS
───────────────
`POST /export/sheets` fills the workbook's monthly sheets with expenses. The
app also owns three things the workbook had nowhere to put: debts, the budget,
and recurring rules. Until they have a home in the sheet, a copy of the
workbook is a record of spending, not a record of the ledger — so if the app
were gone, three quarters of what the owner tracks would be gone with it.

This script adds that home. It is a template migration, not a data tool: it
writes headers, formulas, dropdowns and conditional formats, and never a
figure. The sync fills the rows.

WHAT IT GUARANTEES
──────────────────
* Existing ledger sheets are refused unless --replace-ledger is supplied.
  That option rebuilds them without their records, only in a NEW output file.
  The source and any existing output file are never overwritten.
* The geometry it writes is the geometry `app/routers/sheets.py` writes into:
  tab titles, first and last row, and which columns are the sheet's own.
  Both sides name the same constants; changing one without the other is the
  bug this docstring exists to make obvious.
* Derived columns are formulas and the sync never writes them — `অবস্থা` on
  the debt sheet, `গ্রুপ` and `মাসিক সমমান` on the recurring sheet, and the
  four computed columns on the budget sheet. This is the same division of
  labour as `গ্রুপ` on the monthly sheets: the taxonomy has one source.
* Nothing outside the three new sheets changes, except two additions the new
  sheets depend on — the `MonthTabs` name on `সেটিংস` (the budget sheet's
  month picker reads it) and the guide's sheet inventory.
"""

import argparse
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ── Names and geometry shared with app/routers/sheets.py ────────────────────
TAB_DEBTS = "ধার-দেনা"
TAB_BUDGET = "বাজেট"
TAB_RECURRING = "পুনরাবৃত্ত খরচ"

DEBT_FIRST, DEBT_LAST = 4, 103
BUDGET_FIRST, BUDGET_LAST = 4, 53
RECUR_FIRST, RECUR_LAST = 4, 103
#: Where the sync writes the budget's single total. Not a row — one cell.
BUDGET_TOTAL_CELL = "D2"

GUIDE = "নির্দেশিকা"
YEARLY = "বার্ষিক সারসংক্ষেপ"
SETTINGS = "সেটিংস"

BN_MONTHS = [
    "জানুয়ারি", "ফেব্রুয়ারি", "মার্চ", "এপ্রিল", "মে", "জুন",
    "জুলাই", "আগস্ট", "সেপ্টেম্বর", "অক্টোবর", "নভেম্বর", "ডিসেম্বর",
]
YEAR_BN = "২০২৬"

#: The exact strings the app's enums arrive as. The dropdowns validate against
#: these, so a change here without the matching change in `_DIR_BN` /
#: `_FREQ_BN` / `_ACTIVE_BN` makes every synced row fail validation.
DIR_LEND, DIR_BORROW = "ধার দিয়েছি", "ধার নিয়েছি"
FREQ = ["প্রতিদিন", "প্রতি সপ্তাহে", "প্রতি মাসে", "প্রতি বছরে"]
ACTIVE_ON, ACTIVE_OFF = "চালু", "বন্ধ"
STATE_OPEN, STATE_SETTLED = "চলমান", "পরিশোধিত"

# ── The workbook's own visual vocabulary, read off its monthly sheets ───────
NAVY, BLUE, WHITE = "FF1F3864", "FF2E5C99", "FFFFFFFF"
GREY_BG, MUTED, HINT_FG = "FFF5F5F5", "FF404040", "FF595959"
PANEL = "FFF2F6FB"
RED_BG, RED_FG, AMBER_BG = "FFFFC7CE", "FF9C0006", "FFFFF2CC"
BN, EN = "Nirmala UI", "Arial"

MONEY = '"৳ "#,##0.00;[RED]"-৳ "#,##0.00;\\–'
PCT = '0.0%;\\-0.0%;\\–'
WHOLE = '#,##0;\\-#,##0;\\–'
DATED = "dd/mm/yyyy"

THIN = Side(style="thin", color="FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

report: list[str] = []


def log(message: str) -> None:
    report.append(message)
    print(message)


# ── Cell writers ───────────────────────────────────────────────────────────


def title(ws, span: str, text: str, size: int = 15) -> None:
    ws.merge_cells(span)
    cell = ws[span.split(":")[0]]
    cell.value = text
    cell.font = Font(name=BN, size=size, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def hint(ws, span: str, text: str) -> None:
    ws.merge_cells(span)
    cell = ws[span.split(":")[0]]
    cell.value = text
    cell.font = Font(name=BN, size=9, italic=True, color=HINT_FG)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def header(ws, row: int, first_col: int, labels: list[str]) -> None:
    for offset, label in enumerate(labels):
        cell = ws.cell(row, first_col + offset)
        cell.value = label
        cell.font = Font(name=BN, size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def body(
    ws,
    coord: str,
    *,
    value=None,
    fmt: str = "General",
    font: str = BN,
    align: str = "left",
    derived: bool = False,
    bold: bool = False,
) -> None:
    """An ordinary data cell. ``derived`` marks a formula column the sync
    must not write — grey and italic, exactly as `গ্রুপ` is on the monthly
    sheets, so "do not type here" is visible rather than documented."""
    cell = ws[coord]
    if value is not None:
        cell.value = value
    cell.font = Font(
        name=font, size=11, italic=derived, bold=bold, color=MUTED if derived else "FF000000"
    )
    cell.fill = PatternFill("solid", fgColor=GREY_BG if derived else WHITE)
    cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BOX


def band(ws, coord: str, value=None, *, fmt: str = "General", align: str = "left",
         size: int = 11, italic: bool = False) -> None:
    """A row of the navy summary band (the monthly sheets' `I12:J13` idiom)."""
    cell = ws[coord]
    if value is not None:
        cell.value = value
    cell.font = Font(name=BN, size=size, bold=not italic, italic=italic, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BOX


def panel(ws, coord: str, value=None, *, fmt: str = "General", align: str = "left") -> None:
    """A row of the pale side panel (the monthly sheets' `J16` idiom)."""
    cell = ws[coord]
    if value is not None:
        cell.value = value
    cell.font = Font(name=BN, size=11, bold=True, color=NAVY)
    cell.fill = PatternFill("solid", fgColor=PANEL)
    cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = BOX


def widths(ws, mapping: dict[str, float]) -> None:
    for letter, width in mapping.items():
        ws.column_dimensions[letter].width = width


def heights(ws, first: int, last: int) -> None:
    ws.row_dimensions[1].height = 25.5
    ws.row_dimensions[2].height = 15.75
    ws.row_dimensions[3].height = 21.75
    for row in range(first, last + 1):
        ws.row_dimensions[row].height = 15.0


def chrome(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = NAVY


def listed(ws, span: str, options: list[str] | str, *, stop: bool = True,
           title_bn: str, message_bn: str) -> None:
    """Dropdown validation. A list of options becomes an inline list; a string
    is used as-is so a defined name (`=CategoryList`) can be passed."""
    formula = options if isinstance(options, str) else '"' + ",".join(options) + '"'
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop" if stop else "warning",
        errorTitle=title_bn,
        error=message_bn,
    )
    ws.add_data_validation(validation)
    validation.add(span)


def numeric(ws, span: str, *, title_bn: str, message_bn: str) -> None:
    validation = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle=title_bn,
        error=message_bn,
    )
    ws.add_data_validation(validation)
    validation.add(span)


def dates(ws, span: str) -> None:
    validation = DataValidation(
        type="date",
        operator="greaterThanOrEqual",
        formula1="DATE(2000,1,1)",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="তারিখ দেখে নিন",
        error="এটি তারিখ হিসেবে পড়া যায়নি। DD/MM/YYYY আকারে লিখুন।",
    )
    ws.add_data_validation(validation)
    validation.add(span)


def flag(ws, span: str, condition: str, bg: str, fg: str | None = None) -> None:
    ws.conditional_formatting.add(
        span,
        FormulaRule(
            formula=[condition],
            fill=PatternFill("solid", bgColor=bg),
            font=Font(color=fg) if fg else None,
            stopIfTrue=fg is not None,
        ),
    )


def fresh(wb, name: str, position: int):
    """Create the sheet, replacing any earlier version of it in place."""
    if name in wb.sheetnames:
        position = wb.sheetnames.index(name)
        del wb[name]
        log(f"  rebuilt existing sheet ‘{name}’")
    return wb.create_sheet(name, position)


# ── সেটিংস: the month list the budget sheet's picker reads ──────────────────


def month_tabs(wb) -> None:
    ws = wb[SETTINGS]
    ws["I1"] = "মাসিক শিটের তালিকা"
    ws["I1"].font = Font(name=BN, size=11, bold=True, color=WHITE)
    ws["I1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["I1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["I1"].border = BOX
    for index, month in enumerate(BN_MONTHS):
        cell = ws.cell(2 + index, 9)
        cell.value = f"{month} {YEAR_BN}"
        cell.font = Font(name=BN, size=11, color="FF000000")
        cell.fill = PatternFill("solid", fgColor=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BOX
    widths(ws, {"H": 3.0, "I": 26.0})
    existing = {name for name in wb.defined_names}
    if "MonthTabs" in existing:
        del wb.defined_names["MonthTabs"]
    wb.defined_names.add(
        openpyxl.workbook.defined_name.DefinedName(
            "MonthTabs", attr_text=f"{SETTINGS}!$I$2:$I$13"
        )
    )
    log(f"  {SETTINGS}: I1:I13 month list + defined name MonthTabs")


# ── ধার-দেনা ───────────────────────────────────────────────────────────────


def debts_sheet(wb) -> None:
    ws = fresh(wb, TAB_DEBTS, 2)
    widths(ws, {"A": 13.0, "B": 26.0, "C": 22.0, "D": 15.0, "E": 34.0, "F": 14.0,
                "G": 18.0, "H": 3.0, "I": 34.0, "J": 17.0})
    heights(ws, DEBT_FIRST, DEBT_LAST)
    chrome(ws)

    title(ws, "A1:G1", "ধার-দেনা  —  কে কত পাবে / দেবে")
    title(ws, "I1:J1", "সারসংক্ষেপ", size=13)
    hint(
        ws,
        "A2:G2",
        "অ্যাপ থেকে সিঙ্ক করলে এই শিটের ৪–১০৩ সারি সম্পূর্ণ নতুন করে লেখা হয়। "
        "‘অবস্থা’ কলাম স্বয়ংক্রিয় — এখানে কিছু লিখবেন না। ‘বাকি’ মানে এখনো "
        "অপরিশোধিত অংশ, মূল ধারের পরিমাণ নয়।",
    )
    header(ws, 3, 1, ["তারিখ", "কার সাথে", "ধরন", "বাকি (৳)", "নোট", "অবস্থা",
                      "পরিশোধের তারিখ"])
    header(ws, 3, 9, ["হিসাব", "মোট (৳)"])

    for row in range(DEBT_FIRST, DEBT_LAST + 1):
        body(ws, f"A{row}", fmt=DATED, font=EN, align="center")
        body(ws, f"B{row}")
        body(ws, f"C{row}")
        body(ws, f"D{row}", fmt=MONEY, font=EN, align="right")
        body(ws, f"E{row}")
        body(
            ws,
            f"F{row}",
            value=f'=IF($A{row}="","",IF($G{row}="","{STATE_OPEN}","{STATE_SETTLED}"))',
            derived=True,
            align="center",
        )
        body(ws, f"G{row}", fmt=DATED, font=EN, align="center")

    span = f"$D${DEBT_FIRST}:$D${DEBT_LAST}"
    kind = f"$C${DEBT_FIRST}:$C${DEBT_LAST}"
    paid = f"$G${DEBT_FIRST}:$G${DEBT_LAST}"
    when = f"$A${DEBT_FIRST}:$A${DEBT_LAST}"
    # `""` as a SUMIFS criterion matches empty cells and `"<>"` matches
    # non-empty ones, in both Excel and Google Sheets. The settle date is the
    # single source of "is this still open" — the same rule the `অবস্থা`
    # formula above uses, so the panel and the column can never disagree.
    panel(ws, "I4", f"মোট পাবো ({DIR_LEND}) — {STATE_OPEN}")
    panel(ws, "J4", f'=SUMIFS({span},{kind},"{DIR_LEND}",{paid},"")', fmt=MONEY, align="right")
    panel(ws, "I5", f"মোট দেবো ({DIR_BORROW}) — {STATE_OPEN}")
    panel(ws, "J5", f'=SUMIFS({span},{kind},"{DIR_BORROW}",{paid},"")', fmt=MONEY, align="right")
    band(ws, "I6", "নিট অবস্থান  (পাবো − দেবো)")
    band(ws, "J6", "=$J$4-$J$5", fmt=MONEY, align="right")
    panel(ws, "I7", f"{STATE_SETTLED} (মোট)")
    panel(ws, "J7", f'=SUMIFS({span},{paid},"<>")', fmt=MONEY, align="right")
    panel(ws, "I8", f"{STATE_OPEN} এন্ট্রি")
    panel(ws, "J8", f'=COUNTIFS({when},"<>",{paid},"")', fmt=WHOLE, align="right")
    panel(ws, "I9", "মোট এন্ট্রি")
    panel(ws, "J9", f'=COUNTIF({when},"<>")', fmt=WHOLE, align="right")

    listed(ws, f"C{DEBT_FIRST}:C{DEBT_LAST}", [DIR_LEND, DIR_BORROW],
           title_bn="ধরন বেছে নিন",
           message_bn=f"‘{DIR_LEND}’ বা ‘{DIR_BORROW}’ — ড্রপ-ডাউন থেকে বেছে নিন।")
    numeric(ws, f"D{DEBT_FIRST}:D{DEBT_LAST}", title_bn="পরিমাণ ঠিক নয়",
            message_bn="শুধু শূন্য বা তার বেশি সংখ্যা লিখুন।")
    dates(ws, f"A{DEBT_FIRST}:A{DEBT_LAST}")
    dates(ws, f"G{DEBT_FIRST}:G{DEBT_LAST}")

    rows = f"A{DEBT_FIRST}:G{DEBT_LAST}"
    # An entry with a party and no amount is the debt sheet's version of the
    # bug the monthly sheets' reconciliation line catches: it sits in plain
    # sight and contributes nothing to any total. Red, and checked first.
    flag(ws, rows, f'AND($A{DEBT_FIRST}<>"",$D{DEBT_FIRST}="")', RED_BG, RED_FG)
    flag(ws, rows, f'AND($A{DEBT_FIRST}<>"",$G{DEBT_FIRST}="")', AMBER_BG)
    log(f"  {TAB_DEBTS}: rows {DEBT_FIRST}–{DEBT_LAST}, F derived, panel I4:J9")


# ── বাজেট ──────────────────────────────────────────────────────────────────


def budget_sheet(wb) -> None:
    ws = fresh(wb, TAB_BUDGET, 3)
    widths(ws, {"A": 34.0, "B": 17.0, "C": 19.0, "D": 17.0, "E": 12.0, "F": 20.0})
    heights(ws, BUDGET_FIRST, BUDGET_LAST + 2)
    chrome(ws)

    title(ws, "A1:F1", "বাজেট  —  খরচ বনাম সীমা")
    # The month picker is what makes one budget sheet work for twelve months:
    # every actual-spend formula reads the tab this cell names. Written as a
    # dropdown over the real tab titles so it can never name a tab that is
    # not there.
    band(ws, "A2", "মাস", align="center")
    body(ws, "B2", value=f"{BN_MONTHS[0]} {YEAR_BN}", align="center", bold=True)
    band(ws, "C2", "মোট মাসিক বাজেট (৳)", align="center")
    body(ws, BUDGET_TOTAL_CELL, fmt=MONEY, font=EN, align="right", bold=True)
    ws.merge_cells("E2:F2")
    over = ws["E2"]
    over.value = (
        '=IF($D$2=0,"",IF(SUM($B$4:$B$53)>$D$2,'
        '"⚠ খাতভিত্তিক সীমার যোগফল মোট বাজেটের বেশি",""))'
    )
    over.font = Font(name=BN, size=10, bold=True, color=RED_FG)
    over.alignment = Alignment(horizontal="left", vertical="center")

    header(ws, 3, 1, ["খাত", "মাসিক সীমা (৳)", "এই মাসের খরচ (৳)", "বাকি (৳)",
                      "ব্যবহার", "অবস্থা"])

    spent = '"\'"&$B$2&"\'!$C$4:$C$203"'
    amounts = '"\'"&$B$2&"\'!$E$4:$E$203"'
    for row in range(BUDGET_FIRST, BUDGET_LAST + 1):
        body(ws, f"A{row}")
        body(ws, f"B{row}", fmt=MONEY, font=EN, align="right")
        body(
            ws,
            f"C{row}",
            value=(
                f'=IF($A{row}="","",SUMIF(INDIRECT({spent}),$A{row},'
                f"INDIRECT({amounts})))"
            ),
            fmt=MONEY,
            font=EN,
            align="right",
            derived=True,
        )
        body(ws, f"D{row}", value=f'=IF($A{row}="","",$B{row}-$C{row})',
             fmt=MONEY, font=EN, align="right", derived=True)
        # 0 rather than "" when there is no limit: text compares greater than
        # any number in a spreadsheet, so a blank here would make every
        # unlimited row read as "over budget" to the conditional format.
        body(ws, f"E{row}", value=f'=IF($A{row}="","",IFERROR($C{row}/$B{row},0))',
             fmt=PCT, align="center", derived=True)
        body(
            ws,
            f"F{row}",
            value=(
                f'=IF($A{row}="","",IF($B{row}=0,"—",'
                f'IF($E{row}>1,"বেশি হয়ে গেছে",IF($E{row}>=0.8,"সতর্ক","ভালো"))))'
            ),
            align="center",
            derived=True,
        )

    total = BUDGET_LAST + 1
    band(ws, f"A{total}", "মোট")
    band(ws, f"B{total}", f"=SUM($B${BUDGET_FIRST}:$B${BUDGET_LAST})", fmt=MONEY, align="right")
    band(ws, f"C{total}", f"=SUM($C${BUDGET_FIRST}:$C${BUDGET_LAST})", fmt=MONEY, align="right")
    band(ws, f"D{total}", f"=$B${total}-$C${total}", fmt=MONEY, align="right")
    band(ws, f"E{total}", f'=IFERROR($C${total}/$B${total},0)', fmt=PCT, align="center")
    band(ws, f"F{total}", "")

    # The month's own authoritative total minus everything the budget rows
    # absorbed. Non-zero means real spending sits in categories with no limit
    # set — the budget page's blind spot, named instead of invisible.
    loose = BUDGET_LAST + 2
    band(ws, f"A{loose}", "বাজেটবিহীন খাতে খরচ", size=10, italic=True)
    band(ws, f"B{loose}", "")
    band(
        ws,
        f"C{loose}",
        f'=IFERROR(INDIRECT("\'"&$B$2&"\'!$J$12")-$C${total},"")',
        fmt=MONEY,
        align="right",
    )
    band(ws, f"D{loose}", "")
    band(ws, f"E{loose}", "")
    band(ws, f"F{loose}", "")

    listed(ws, "B2", "=MonthTabs", title_bn="মাস বেছে নিন",
           message_bn="ড্রপ-ডাউন থেকে একটি মাসিক শিটের নাম বেছে নিন।")
    listed(ws, f"A{BUDGET_FIRST}:A{BUDGET_LAST}", "=CategoryList",
           title_bn="খাত তালিকায় নেই",
           message_bn="‘সেটিংস’ শিটের তালিকা থেকে খাত বেছে নিন।")
    numeric(ws, f"B{BUDGET_FIRST}:B{BUDGET_LAST}", title_bn="সীমা ঠিক নয়",
            message_bn="শুধু শূন্য বা তার বেশি সংখ্যা লিখুন।")
    numeric(ws, BUDGET_TOTAL_CELL, title_bn="বাজেট ঠিক নয়",
            message_bn="শুধু শূন্য বা তার বেশি সংখ্যা লিখুন।")

    rows = f"A{BUDGET_FIRST}:F{BUDGET_LAST}"
    flag(ws, rows, f'AND($A{BUDGET_FIRST}<>"",$B{BUDGET_FIRST}>0,$E{BUDGET_FIRST}>1)',
         RED_BG, RED_FG)
    flag(ws, rows, f'AND($A{BUDGET_FIRST}<>"",$B{BUDGET_FIRST}>0,$E{BUDGET_FIRST}>=0.8)',
         AMBER_BG)
    log(
        f"  {TAB_BUDGET}: rows {BUDGET_FIRST}–{BUDGET_LAST}, total in "
        f"{BUDGET_TOTAL_CELL}, C:F derived, reconciliation at row {loose}"
    )


# ── পুনরাবৃত্ত খরচ ──────────────────────────────────────────────────────────


def recurring_sheet(wb) -> None:
    ws = fresh(wb, TAB_RECURRING, 4)
    widths(ws, {"A": 34.0, "B": 24.0, "C": 15.0, "D": 22.0, "E": 30.0, "F": 18.0,
                "G": 15.0, "H": 15.0, "I": 11.0, "J": 19.0, "K": 3.0, "L": 32.0,
                "M": 17.0})
    heights(ws, RECUR_FIRST, RECUR_LAST)
    chrome(ws)

    title(ws, "A1:J1", "পুনরাবৃত্ত খরচ  —  নিয়মিত বিল ও কিস্তি")
    title(ws, "L1:M1", "সারসংক্ষেপ", size=13)
    hint(
        ws,
        "A2:J2",
        "অ্যাপ থেকে সিঙ্ক করলে এই শিটের ৪–১০৩ সারি সম্পূর্ণ নতুন করে লেখা হয়। "
        "‘গ্রুপ’ ও ‘মাসিক সমমান’ কলাম স্বয়ংক্রিয় — এখানে কিছু লিখবেন না। "
        "এগুলো নিয়ম, খরচ নয়: প্রকৃত খরচ মাসিক শিটে যায়।",
    )
    header(ws, 3, 1, ["খাত", "গ্রুপ", "পরিমাণ (৳)", "পেমেন্ট মাধ্যম", "বিবরণ",
                      "কত দিন পর পর", "শুরুর তারিখ", "পরবর্তী", "চালু?",
                      "মাসিক সমমান (৳)"])
    header(ws, 3, 12, ["হিসাব", "মান"])

    for row in range(RECUR_FIRST, RECUR_LAST + 1):
        body(ws, f"A{row}")
        # The same lookup the monthly sheets use in column D, so a rule and an
        # expense in the same category can never land in different groups.
        body(ws, f"B{row}",
             value=f'=IFERROR(INDEX(GroupLookup,MATCH($A{row},CategoryList,0)),"")',
             derived=True)
        body(ws, f"C{row}", fmt=MONEY, font=EN, align="right")
        body(ws, f"D{row}")
        body(ws, f"E{row}")
        body(ws, f"F{row}", align="center")
        body(ws, f"G{row}", fmt=DATED, font=EN, align="center")
        body(ws, f"H{row}", fmt=DATED, font=EN, align="center")
        body(ws, f"I{row}", align="center")
        # Nested IFs, not an array constant: `{30;4.33;1}` means different
        # things in different spreadsheet locales, and this file is opened in
        # both Excel and Google Sheets.
        body(
            ws,
            f"J{row}",
            value=(
                f'=IF($A{row}="","",IF($I{row}<>"{ACTIVE_ON}",0,$C{row}*'
                f'IF($F{row}="{FREQ[0]}",30,IF($F{row}="{FREQ[1]}",4.33,'
                f'IF($F{row}="{FREQ[2]}",1,IF($F{row}="{FREQ[3]}",1/12,0))))))'
            ),
            fmt=MONEY,
            font=EN,
            align="right",
            derived=True,
        )

    live = f"$I${RECUR_FIRST}:$I${RECUR_LAST}"
    due = f"$H${RECUR_FIRST}:$H${RECUR_LAST}"
    panel(ws, "L4", f"{ACTIVE_ON} নিয়ম")
    panel(ws, "M4", f'=COUNTIF({live},"{ACTIVE_ON}")', fmt=WHOLE, align="right")
    panel(ws, "L5", f"{ACTIVE_OFF} নিয়ম")
    panel(ws, "M5", f'=COUNTIF({live},"{ACTIVE_OFF}")', fmt=WHOLE, align="right")
    band(ws, "L6", "মোট মাসিক প্রতিশ্রুতি")
    band(ws, "M6", f"=SUM($J${RECUR_FIRST}:$J${RECUR_LAST})", fmt=MONEY, align="right")
    band(ws, "L7", "মোট বার্ষিক প্রতিশ্রুতি")
    band(ws, "M7", "=$M$6*12", fmt=MONEY, align="right")
    # A rule whose next date has passed is one the app has not materialized —
    # the one thing about a recurring rule that goes wrong silently.
    panel(ws, "L8", "পরবর্তী তারিখ পেরিয়ে গেছে")
    panel(ws, "M8", f'=COUNTIFS({live},"{ACTIVE_ON}",{due},"<"&TODAY())',
          fmt=WHOLE, align="right")

    listed(ws, f"A{RECUR_FIRST}:A{RECUR_LAST}", "=CategoryList",
           title_bn="খাত তালিকায় নেই",
           message_bn="‘সেটিংস’ শিটের তালিকা থেকে খাত বেছে নিন।")
    listed(ws, f"D{RECUR_FIRST}:D{RECUR_LAST}", "=PaymentList",
           title_bn="পেমেন্ট মাধ্যম বেছে নিন",
           message_bn="‘সেটিংস’ শিটের তালিকা থেকে পেমেন্ট মাধ্যম বেছে নিন।")
    listed(ws, f"F{RECUR_FIRST}:F{RECUR_LAST}", FREQ,
           title_bn="পুনরাবৃত্তি বেছে নিন",
           message_bn="প্রতিদিন / প্রতি সপ্তাহে / প্রতি মাসে / প্রতি বছরে — একটি বেছে নিন।")
    listed(ws, f"I{RECUR_FIRST}:I{RECUR_LAST}", [ACTIVE_ON, ACTIVE_OFF],
           title_bn="অবস্থা বেছে নিন",
           message_bn=f"‘{ACTIVE_ON}’ বা ‘{ACTIVE_OFF}’ বেছে নিন।")
    numeric(ws, f"C{RECUR_FIRST}:C{RECUR_LAST}", title_bn="পরিমাণ ঠিক নয়",
            message_bn="শুধু শূন্য বা তার বেশি সংখ্যা লিখুন।")
    dates(ws, f"G{RECUR_FIRST}:G{RECUR_LAST}")
    dates(ws, f"H{RECUR_FIRST}:H{RECUR_LAST}")

    rows = f"A{RECUR_FIRST}:J{RECUR_LAST}"
    flag(ws, rows, f'AND($A{RECUR_FIRST}<>"",$C{RECUR_FIRST}="")', RED_BG, RED_FG)
    flag(
        ws,
        rows,
        f'AND($I{RECUR_FIRST}="{ACTIVE_ON}",$H{RECUR_FIRST}<>"",'
        f"$H{RECUR_FIRST}<TODAY())",
        AMBER_BG,
    )
    log(f"  {TAB_RECURRING}: rows {RECUR_FIRST}–{RECUR_LAST}, B and J derived, panel L4:M8")


# ── নির্দেশিকা: keep the guide's own inventory true ─────────────────────────


def guide(wb) -> None:
    ws = wb[GUIDE]
    ws["C4"] = (
        "১৮টি শিট: এই নির্দেশিকা, একটি বার্ষিক সারসংক্ষেপ, তিনটি লেজার শিট "
        "(ধার-দেনা, বাজেট, পুনরাবৃত্ত খরচ), ১২টি মাসিক শিট (জানুয়ারি–ডিসেম্বর ২০২৬) "
        "এবং একটি সেটিংস শিট।"
    )
    ws["C5"] = (
        "মাসিক শিটের A থেকে G কলামে (৪–২০৩ সারি), এবং তিনটি লেজার শিটে। "
        "ধূসর ও ইটালিক ঘরগুলো সূত্র — সেখানে কিছু লিখবেন না।"
    )
    entries = [
        (
            "ধার-দেনা শিট",
            "কে কত পাবে, কে কত দেবে। ‘বাকি’ কলামে এখনো অপরিশোধিত অংশ থাকে — "
            "আংশিক পরিশোধ হলে অঙ্কটি কমে যায়, মূল ধারের পরিমাণ আলাদা করে রাখা হয় না। "
            "‘পরিশোধের তারিখ’ ফাঁকা থাকলে ধারটি চলমান; ডান পাশের সারসংক্ষেপে "
            "পাবো, দেবো ও নিট অবস্থান আলাদা করে দেখা যায়।",
        ),
        (
            "বাজেট শিট",
            "উপরের ‘মাস’ ড্রপ-ডাউন যে মাসিক শিটকে দেখায়, সেই মাসের প্রকৃত খরচ "
            "খাত ধরে ধরে সীমার পাশে বসে। ৮০% ছুঁলে হলুদ, সীমা ছাড়ালে লাল। "
            "একদম নিচের ‘বাজেটবিহীন খাতে খরচ’ ঘরটি শূন্য না হলে বোঝা যাবে কোনো "
            "খাতে খরচ হচ্ছে যার কোনো সীমা ঠিক করা নেই।",
        ),
        (
            "পুনরাবৃত্ত খরচ শিট",
            "বাসা ভাড়া, বিদ্যুৎ বিলের মতো নিয়মিত খরচের নিয়ম — খরচ নয়, নিয়ম। "
            "‘মাসিক সমমান’ কলাম প্রতিটি নিয়মকে মাসিক অঙ্কে দাঁড় করায় (দৈনিক ×৩০, "
            "সাপ্তাহিক ×৪.৩৩, বার্ষিক ÷১২), তাই ডান পাশে মোট মাসিক ও বার্ষিক "
            "প্রতিশ্রুতি এক নজরে দেখা যায়।",
        ),
        (
            "অ্যাপ থেকে সিঙ্ক",
            "‘পই পই হিসাব’ অ্যাপের সেটিংস থেকে সিঙ্ক করলে মাসিক শিটের খরচ এবং "
            "এই তিনটি লেজার শিট — তিনটিই সম্পূর্ণ নতুন করে লেখা হয়। একই সিঙ্ক "
            "বারবার করলে কিছু দ্বিগুণ হয় না। শিটে হাতে করা বদল পরের সিঙ্কে মুছে "
            "যাবে, তাই অ্যাপ চালু থাকলে বদলগুলো অ্যাপে করুন।",
        ),
    ]
    row = 23
    for label, text in entries:
        ws.cell(row, 2).value = label
        ws.cell(row, 2)._style = ws.cell(17, 2)._style
        ws.cell(row, 3).value = text
        ws.cell(row, 3)._style = ws.cell(17, 3)._style
        ws.row_dimensions[row].height = 45.0
        row += 1
    log(f"  {GUIDE}: sheet count corrected to ১৮, four rows added (B23:C26)")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Add ledger tabs to a template copy.")
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path, help="New output file; must not exist")
    parser.add_argument(
        "--replace-ledger", action="store_true",
        help="Discard existing ledger tabs and their records in the NEW output copy",
    )
    args = parser.parse_args(argv)
    source, target = args.source.resolve(), args.target.resolve()
    if source == target or target.exists():
        print("error: choose a new output file; source and existing files are never overwritten")
        return 1
    wb = openpyxl.load_workbook(source)
    for required in (GUIDE, YEARLY, SETTINGS):
        if required not in wb.sheetnames:
            print(f"error: ‘{required}’ is missing — is this the right workbook?")
            return 1

    existing = [name for name in (TAB_DEBTS, TAB_BUDGET, TAB_RECURRING) if name in wb]
    if existing and not args.replace_ledger:
        print("error: existing ledger tabs contain potentially valuable records: " + ", ".join(existing))
        print("Use --replace-ledger only to rebuild a template copy without those records.")
        return 1

    log(f"reading {source}")
    month_tabs(wb)
    debts_sheet(wb)
    budget_sheet(wb)
    recurring_sheet(wb)
    guide(wb)
    # Exclusive creation also refuses a file created after the initial check.
    with target.open("xb") as output:
        wb.save(output)
    log(f"wrote {target} — {len(wb.sheetnames)} sheets")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
